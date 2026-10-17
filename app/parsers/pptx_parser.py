"""PowerPoint（.pptx）解析器。

选型：python-pptx
关键点：**备注页（notes）也要提取**——企业内部 PPT 的关键结论、口径说明
经常写在备注里，只读正文会丢掉最重要的信息。
"""
import io
import logging

from pptx import Presentation

from app.config import settings
from app.parsers.base import Block, ParsedDocument, rows_to_markdown

logger = logging.getLogger(__name__)


class PptxParser:
    """PPT 解析：标题 + 正文 + 表格 + 备注。"""

    name = "pptx"
    extensions = (".pptx",)

    def parse(self, content: bytes, filename: str) -> ParsedDocument:
        prs = Presentation(io.BytesIO(content))
        blocks: list[Block] = []
        warnings: list[str] = []
        notes_count = table_count = 0

        for idx, slide in enumerate(prs.slides, start=1):
            title = self._slide_title(slide)
            if title:
                blocks.append(Block(type="title", text=title, level=1, page=idx))

            for shape in slide.shapes:
                if shape.has_table:
                    table_count += 1
                    rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
                    md, semantic, warns = rows_to_markdown(
                        rows,
                        table_name=title or f"第 {idx} 页表格",
                        max_cols=settings.TABLE_MAX_COLS,
                        max_rows=settings.TABLE_MAX_ROWS,
                    )
                    warnings.extend(f"[第 {idx} 页] {w}" for w in warns)
                    if md:
                        blocks.append(Block(type="table", text=md, page=idx, meta={"semantic": semantic}))
                elif shape.has_text_frame:
                    text = "\n".join(
                        (p.text or "").strip() for p in shape.text_frame.paragraphs
                    ).strip()
                    if not text or text == title:
                        continue
                    blocks.append(Block(type="paragraph", text=text, page=idx))

            # 备注页
            if slide.has_notes_slide:
                notes = (slide.notes_slide.notes_text_frame.text or "").strip()
                if notes:
                    notes_count += 1
                    blocks.append(
                        Block(
                            type="paragraph",
                            text=f"[备注] {notes}",
                            page=idx,
                            meta={"notes": True},
                        )
                    )

        if not blocks:
            warnings.append("PPT 中未解析到文本内容")

        text = "\n".join(b.retrieval_text for b in blocks)
        return ParsedDocument(
            text=text,
            blocks=blocks,
            metadata={
                "format": "pptx",
                "slides": len(prs.slides),
                "tables": table_count,
                "notes": notes_count,
                "block_count": len(blocks),
            },
            warnings=warnings,
        )

    @staticmethod
    def _slide_title(slide) -> str:
        try:
            if slide.shapes.title is not None:
                return (slide.shapes.title.text or "").strip()
        except Exception:  # noqa: BLE001
            pass
        return ""


__all__ = ["PptxParser"]
