"""解析子系统单元测试（全离线：OCR 打桩，不联网、不加载模型）。"""
import io
import json

import pytest

from app.parsers import (
    UnsupportedFormatError,
    content_md5,
    parse_document,
    parser_for,
    supported_extensions,
)
from app.parsers.base import ParsedDocument, rows_to_markdown


# ======================================================================
# 测试样本构造
# ======================================================================
def build_docx() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading("第三章 休假制度", level=1)
    doc.add_paragraph("员工入职满一年后，每年享有5天带薪年假；工作满三年后，年假增加至10天。")
    doc.add_heading("3.1 差旅住宿标准", level=2)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "城市级别"
    table.cell(0, 1).text = "住宿标准"
    table.cell(1, 0).text = "一线城市"
    table.cell(1, 1).text = "500 元/晚"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_xlsx() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "库存表"
    ws.append(["产品", "库存", "单价"])
    ws.append(["产品A", 120, 35])
    ws.append(["产品B", 80, 42])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_pptx() -> bytes:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
    slide.shapes.title.text = "第三季度经营总结"
    notes = slide.notes_slide.notes_text_frame
    notes.text = "关键结论：营收环比增长20%，主要来自华东区。"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def build_scanned_pdf() -> bytes:
    """生成无文本层的 PDF（仅一张图）——用于扫描件判定测试。"""
    import numpy as np
    import pymupdf
    from PIL import Image

    img = Image.fromarray(np.full((200, 400, 3), 255, dtype="uint8"))
    img_bytes = io.BytesIO()
    img.save(img_bytes, format="PNG")

    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_image(pymupdf.Rect(0, 0, 400, 200), stream=img_bytes.getvalue())
    out = doc.tobytes()
    doc.close()
    return out


# ======================================================================
# 注册表与分发
# ======================================================================
class TestRegistry:
    def test_supported_extensions_cover_office_and_ocr(self):
        exts = supported_extensions()
        for ext in (".pdf", ".docx", ".xlsx", ".pptx", ".csv", ".html", ".json", ".txt", ".md", ".png"):
            assert ext in exts, f"缺少格式支持：{ext}"

    def test_parser_lookup(self):
        assert parser_for("a.DOCX").name == "docx"   # 大小写不敏感
        assert parser_for("报表.xlsx").name == "xlsx"
        assert parser_for("扫描件.pdf").name == "pdf"

    def test_unsupported_format(self):
        with pytest.raises(UnsupportedFormatError):
            parser_for("old.doc")
        with pytest.raises(UnsupportedFormatError):
            parse_document(b"x", "noext")


# ======================================================================
# 各格式解析
# ======================================================================
class TestDocx:
    def test_structure_and_table(self):
        parsed, cached = parse_document(build_docx(), "员工手册.docx")
        assert cached is False
        assert parsed.metadata["format"] == "docx"
        titles = [b for b in parsed.blocks if b.type == "title"]
        assert [t.text for t in titles] == ["第三章 休假制度", "3.1 差旅住宿标准"]
        assert titles[1].level == 2
        tables = [b for b in parsed.blocks if b.type == "table"]
        assert len(tables) == 1
        assert "一线城市" in tables[0].text
        # 表格语义化（检索用）
        semantic = tables[0].meta["semantic"]
        assert any("住宿标准 为 500 元/晚" in s for s in semantic)


class TestExcel:
    def test_sheet_block_and_semantic(self):
        parsed, _ = parse_document(build_xlsx(), "库存.xlsx")
        assert parsed.metadata["sheet_count"] == 1
        table = next(b for b in parsed.blocks if b.type == "table")
        assert table.sheet == "库存表"
        # 「产品A 库存 120」应被还原为自然语言句
        assert any("库存 为 120" in s for s in table.meta["semantic"])


class TestPptx:
    def test_title_and_notes(self):
        parsed, _ = parse_document(build_pptx(), "季度总结.pptx")
        assert parsed.metadata["slides"] == 1
        assert any(b.text == "第三季度经营总结" for b in parsed.blocks)
        # 备注页必须被提取（关键结论常写在备注）
        assert parsed.metadata["notes"] == 1
        assert any("营收环比增长20%" in b.text for b in parsed.blocks)


class TestTextLike:
    def test_markdown_headings(self):
        content = "# 总则\n本手册适用于全体员工。\n## 考勤\n每日9:00上班。\n".encode("utf-8")
        parsed, _ = parse_document(content, "制度.md")
        titles = [b for b in parsed.blocks if b.type == "title"]
        assert [(t.text, t.level) for t in titles] == [("总则", 1), ("考勤", 2)]

    def test_markdown_ordered_list_not_heading(self):
        """INC-09 回归：Markdown 有序列表不能被当成章节标题（否则首项内容会丢）。"""
        content = (
            "## 部署流程\n"
            "1. 上传发布包至 /opt/app/release 目录；\n"
            "2. 执行 ./deploy.sh 完成依赖安装；\n"
        ).encode("utf-8")
        parsed, _ = parse_document(content, "运维.md", use_cache=False)
        titles = [b.text for b in parsed.blocks if b.type == "title"]
        assert titles == ["部署流程"]
        assert not any(t.startswith("1.") for t in titles)

    def test_plain_text_manual_heading(self):
        """纯文本仍保留手工标题识别，但带列表标点结尾的不算标题。"""
        content = "第一章 总则\n本手册适用于全体员工。\n1. 上传发布包至 /opt/app/release 目录；\n".encode("utf-8")
        parsed, _ = parse_document(content, "手册.txt", use_cache=False)
        titles = [b.text for b in parsed.blocks if b.type == "title"]
        assert "第一章 总则" in titles
        assert not any("上传发布包" in t for t in titles)

    def test_gbk_fallback(self):
        parsed, _ = parse_document("年假规定：5天。".encode("gbk"), "旧文件.txt")
        assert "年假" in parsed.text
        assert parsed.metadata["encoding"].lower() in {"gbk", "gb18030"}

    def test_json_flatten(self):
        content = json.dumps({"policy": {"annual_leave": 5}}, ensure_ascii=False).encode("utf-8")
        parsed, _ = parse_document(content, "config.json")
        assert "annual_leave" in parsed.text
        assert any("annual_leave 为 5" in s for s in parsed.blocks[0].meta["semantic"])

    def test_csv(self):
        parsed, _ = parse_document("产品,库存\n产品A,120\n".encode("utf-8"), "库存.csv")
        table = parsed.blocks[0]
        assert table.type == "table"
        assert any("库存 为 120" in s for s in table.meta["semantic"])

    def test_html(self):
        html = "<html><body><h1>员工手册</h1><p>总则内容</p><table><tr><th>项</th><th>值</th></tr><tr><td>年假</td><td>5天</td></tr></table></body></html>"
        parsed, _ = parse_document(html.encode("utf-8"), "a.html")
        assert any(b.type == "title" and b.text == "员工手册" for b in parsed.blocks)
        table = next(b for b in parsed.blocks if b.type == "table")
        assert any("值 为 5天" in s for s in table.meta["semantic"])


# ======================================================================
# 扫描件判定（OCR 打桩，避免测试加载模型）
# ======================================================================
class TestScannedPdf:
    def test_probe_detects_scanned(self):
        content = build_scanned_pdf()
        probe = parser_for("scan.pdf").probe(content)
        assert probe["pages"] == 1
        assert probe["is_scanned"] is True
        assert parser_for("scan.pdf").needs_async(content, "scan.pdf") in (True, False)  # 取决于 OCR 开关

    def test_ocr_path_with_stub(self, monkeypatch):
        from app.parsers import ocr as ocr_mod

        monkeypatch.setattr(ocr_mod, "ocr_available", lambda: True)
        monkeypatch.setattr(
            ocr_mod, "recognize", lambda image, min_conf=0.6: ("扫描件识别文本：年假为5天", [])
        )
        parsed, _ = parse_document(build_scanned_pdf(), "scan.pdf")
        assert parsed.metadata["ocr_used"] is True
        assert "年假为5天" in parsed.text
        assert all(b.meta.get("ocr") for b in parsed.blocks)


# ======================================================================
# 解析缓存
# ======================================================================
class TestCache:
    def test_second_parse_hits_cache(self):
        content = build_xlsx()
        md5 = content_md5(content)
        first, cached1 = parse_document(content, "库存.xlsx", md5=md5)
        second, cached2 = parse_document(content, "库存.xlsx", md5=md5)
        assert cached1 is False and cached2 is True
        assert second.text == first.text
        assert isinstance(second, ParsedDocument)

    def test_cache_disabled(self):
        content = "# 标题\n内容".encode("utf-8")
        _, c1 = parse_document(content, "a.md", md5=content_md5(content), use_cache=False)
        _, c2 = parse_document(content, "a.md", md5=content_md5(content), use_cache=False)
        assert c1 is False and c2 is False


# ======================================================================
# 表格语义化工具
# ======================================================================
class TestRowsToMarkdown:
    def test_basic(self):
        md, semantic, warnings = rows_to_markdown(
            [["城市", "标准"], ["一线", "500"]], table_name="差旅"
        )
        assert "| 城市 | 标准 |" in md
        assert semantic == ["差旅：城市 为 一线，标准 为 500。"]
        assert warnings == []

    def test_truncation_warns(self):
        rows = [["列1", "列2"]] + [[str(i), str(i)] for i in range(500)]
        _, _, warnings = rows_to_markdown(rows, max_rows=100)
        assert any("已截断" in w for w in warnings)

    def test_empty(self):
        md, semantic, warnings = rows_to_markdown([])
        assert md == "" and semantic == []
