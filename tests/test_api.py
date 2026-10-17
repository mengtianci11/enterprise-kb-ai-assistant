"""API 层集成测试（TestClient，全离线）。

覆盖 v3.0 新增能力：多格式上传、结构块、异步任务（OCR 打桩）、错误语义。
"""
import io
import time


def _docx_bytes() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading("第三章 休假制度", level=1)
    doc.add_paragraph("员工入职满一年后，每年享有5天带薪年假。")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _xlsx_bytes() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "库存表"
    ws.append(["产品", "库存"])
    ws.append(["产品A", 120])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestUploadAndList:
    def test_upload_txt(self, api_client):
        resp = api_client.post(
            "/documents/upload",
            files={"file": ("制度.txt", "考勤制度：每天工作时间9点到18点。".encode("utf-8"), "text/plain")},
        )
        assert resp.status_code == 200, resp.text
        data = resp.json()
        assert data["status"] == "ready"
        assert data["chunks"] >= 1
        assert data["total_chunks"] >= 1
        assert data["format"] == "text"
        assert data["task_id"] is None
        assert data["blocks"] >= 1

    def test_upload_docx_keeps_structure(self, api_client):
        resp = api_client.post(
            "/documents/upload",
            files={"file": ("员工手册.docx", _docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
        assert resp.status_code == 200, resp.text
        data = resp.json()
        assert data["format"] == "docx"
        assert data["blocks"] >= 2

    def test_upload_xlsx(self, api_client):
        resp = api_client.post(
            "/documents/upload",
            files={"file": ("库存.xlsx", _xlsx_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert resp.status_code == 200, resp.text
        assert resp.json()["format"] == "xlsx"

    def test_upload_unsupported_type(self, api_client):
        resp = api_client.post(
            "/documents/upload",
            files={"file": ("老文档.doc", b"MZ...", "application/octet-stream")},
        )
        assert resp.status_code == 400
        assert "不支持" in resp.json()["detail"]

    def test_malformed_file_returns_400_not_502(self, api_client):
        """损坏文件属客户端问题，应返回 400 并带可读提示（而非 502 服务错误）。"""
        resp = api_client.post(
            "/documents/upload",
            files={"file": ("坏文件.docx", b"this is not a zip file", "application/octet-stream")},
        )
        assert resp.status_code == 400
        assert "解析失败" in resp.json()["detail"]

    def test_upload_duplicate_rejected(self, api_client):
        payload = {"file": ("重复.txt", "相同内容".encode("utf-8"), "text/plain")}
        assert api_client.post("/documents/upload", files=payload).status_code == 200
        resp = api_client.post("/documents/upload", files=payload)
        assert resp.status_code == 409
        assert "已存在" in resp.json()["detail"]

    def test_list_documents_exposes_format(self, api_client):
        api_client.post(
            "/documents/upload",
            files={"file": ("库存.xlsx", _xlsx_bytes(), "application/octet-stream")},
        )
        resp = api_client.get("/documents")
        assert resp.status_code == 200
        doc = resp.json()["documents"][0]
        assert doc["name"] == "库存.xlsx"
        assert doc["format"] == "xlsx"
        assert doc["status"] == "ready"

    def test_delete_document(self, api_client):
        api_client.post(
            "/documents/upload",
            files={"file": ("临时.txt", "临时内容".encode("utf-8"), "text/plain")},
        )
        resp = api_client.delete("/documents/临时.txt")
        assert resp.status_code == 200 and resp.json()["deleted"] is True
        assert api_client.get("/documents").json()["documents"] == []
        assert api_client.delete("/documents/不存在.txt").status_code == 404


class TestAsyncTasks:
    def test_image_upload_goes_async_and_completes(self, api_client, monkeypatch):
        """图片走异步 OCR 路径：立即返回 task_id，随后可轮询到 ready。"""
        from app.parsers import ocr as ocr_mod

        monkeypatch.setattr(ocr_mod, "ocr_available", lambda: True)
        monkeypatch.setattr(
            ocr_mod, "recognize", lambda image, min_conf=0.6: ("扫描图内容：年假为5天", [])
        )

        resp = api_client.post(
            "/documents/upload",
            files={"file": ("凭证.png", b"\x89PNG\r\n\x1a\nfake", "image/png")},
        )
        assert resp.status_code == 200, resp.text
        data = resp.json()
        assert data["status"] == "processing"
        assert data["task_id"]

        task_id = data["task_id"]
        deadline = time.time() + 20
        task = None
        while time.time() < deadline:
            task = api_client.get(f"/documents/tasks/{task_id}").json()
            if task["status"] in {"ready", "failed"}:
                break
            time.sleep(0.2)

        assert task is not None
        assert task["status"] == "ready", task
        assert task["chunks"] >= 1
        # 文档已入库
        assert any(d["name"] == "凭证.png" for d in api_client.get("/documents").json()["documents"])

    def test_task_not_found(self, api_client):
        assert api_client.get("/documents/tasks/不存在").status_code == 404


class TestChat:
    def test_chat_without_docs_returns_400(self, api_client):
        resp = api_client.post("/chat", json={"question": "你好"})
        assert resp.status_code == 400
        assert "上传文档" in resp.json()["detail"]

    def test_chat_with_docs(self, api_client, monkeypatch):
        api_client.post(
            "/documents/upload",
            files={"file": ("制度.txt", "年假规定：入职满一年享5天带薪年假。".encode("utf-8"), "text/plain")},
        )
        from tests.conftest import FakeAgent

        monkeypatch.setattr(api_client.state_mod, "agent", FakeAgent())
        resp = api_client.post("/chat", json={"question": "年假几天", "conversation_id": "t1"})
        assert resp.status_code == 200
        assert "测试回答" in resp.json()["answer"]
        assert resp.json()["request_id"]

    def test_stream_chat(self, api_client, monkeypatch):
        from tests.conftest import FakeAgent

        api_client.post(
            "/documents/upload",
            files={"file": ("制度.txt", "年假规定：入职满一年享5天带薪年假。".encode("utf-8"), "text/plain")},
        )
        monkeypatch.setattr(api_client.state_mod, "agent", FakeAgent())
        with api_client.stream("POST", "/chat/stream", json={"question": "你好"}) as resp:
            assert resp.status_code == 200
            body = "".join(resp.iter_text())
        assert all(ch in body for ch in "测试流式")

    def test_health_reports_parsers_and_ocr(self, api_client):
        resp = api_client.get("/health")
        assert resp.status_code == 200
        data = resp.json()
        assert data["status"] == "ok"
        # 解析器能力可见（运维排查用）
        assert "pdf" in data["parsers"] and "docx" in data["parsers"] and "xlsx" in data["parsers"]
        assert "ocr_available" in data
        assert "pending_tasks" in data
        # 安全：公开的健康检查不再暴露配置详情（已移到需鉴权的 /admin/config）
        assert "config" not in data

    def test_root_serves_frontend(self, api_client):
        resp = api_client.get("/")
        assert resp.status_code == 200
        assert "企业知识库" in resp.text or "助手" in resp.text
